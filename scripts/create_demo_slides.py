from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "demo-slides.pptx"
BLUE = RGBColor(11, 45, 92)
CYAN = RGBColor(0, 174, 239)
RED = RGBColor(225, 32, 52)
INK = RGBColor(20, 30, 48)
MUTED = RGBColor(92, 107, 126)
WHITE = RGBColor(255, 255, 255)
PALE = RGBColor(237, 246, 252)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def textbox(slide, x, y, w, h, text, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    from pptx.enum.shapes import MSO_SHAPE
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def base(title, kicker, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 0.11, RED)
    textbox(slide, .55, .35, 7.6, .3, kicker.upper(), 10, CYAN, True)
    textbox(slide, .55, .68, 12.1, .7, title, 30, BLUE, True)
    textbox(slide, .55, 7.18, 5, .2, "kingpro · AI20K K4 · E403/C1", 8, MUTED)
    textbox(slide, 12.35, 7.14, .4, .25, str(number), 9, BLUE, True, PP_ALIGN.RIGHT)
    return slide


# 1
s = base("Tutor hỏi lại trước khi trả lời câu mơ hồ", "Track A · VLearn Tutor", 1)
textbox(s, .6, 1.55, 7.0, 1.55, "Một học viên hỏi chưa đủ rõ → LangGraph dừng → học viên xác nhận ý định → Tutor mới trả lời.", 27, INK, True)
rect(s, 8.1, 1.5, 4.55, 3.8, BLUE, BLUE, True)
textbox(s, 8.55, 1.95, 3.7, .35, "MỘT VẤN ĐỀ", 11, CYAN, True)
textbox(s, 8.55, 2.48, 3.55, 1.6, "Tutor đang đoán ý và trả lời trước khi biết học viên hỏi gì.", 24, WHITE, True)
textbox(s, .6, 5.05, 7.0, .7, "Demo case: “cho tôi link” trong Bài 16 · Mini Hackathon", 18, BLUE, True)
textbox(s, .6, 5.85, 7.0, .55, "Gemini 3.5 Flash Lite · LangGraph interrupt/resume · key chỉ ở environment", 14, MUTED)

# 2
s = base("Pain có bằng chứng, không phải cảm giác", "Evidence", 2)
stats = [("3.097", "lượt Tutor K4"), ("6", "lượt hỏi probing"), ("2/10", "case live hỏi lại đúng"), ("8/10", "case live trả lời theo giả định")]
for i, (value, label) in enumerate(stats):
    x = .6 + i * 3.15
    rect(s, x, 1.55, 2.75, 1.55, PALE, RGBColor(195, 218, 235), True)
    textbox(s, x+.15, 1.78, 2.45, .55, value, 29, BLUE, True, PP_ALIGN.CENTER)
    textbox(s, x+.18, 2.45, 2.4, .35, label, 11, MUTED, False, PP_ALIGN.CENTER)
textbox(s, .65, 3.55, 12.0, .45, "Lỗi live có thể kiểm chứng", 18, BLUE, True)
rect(s, .65, 4.1, 12.0, 1.45, RGBColor(255, 248, 229), RGBColor(238, 196, 84), True)
textbox(s, .95, 4.30, 11.4, 1.0, "Fail: “cho tôi link” · “form đâu” · “repo nào” · “đáp án gì” · “cái này là gì” · “giải thích đoạn trên” · “cgi” · “giúp tôi với”. Riêng “cho tôi link” trả nhầm repo lớp 3A trên trang 3B.", 16, INK)
textbox(s, .65, 5.9, 12, .5, "Nguồn: tutor_turns.csv + kiểm thử VLearn 17/09 · phương pháp tái lập nằm trong evidence/", 12, MUTED)

# 3
s = base("Một quyết định AI, một điểm dừng có kiểm soát", "LangGraph flow", 3)
nodes = [("1", "Nhận câu hỏi"), ("2", "Phân loại độ rõ"), ("3", "interrupt"), ("4", "Người dùng chọn"), ("5", "resume"), ("6", "Trả lời")]
for i, (n, label) in enumerate(nodes):
    x = .55 + i * 2.08
    rect(s, x, 2.05, 1.72, 1.1, PALE if i not in (2,3) else RGBColor(255,244,207), BLUE if i not in (2,3) else RGBColor(211,154,18), True)
    textbox(s, x+.1, 2.25, 1.52, .25, n, 10, RED, True, PP_ALIGN.CENTER)
    textbox(s, x+.1, 2.58, 1.52, .35, label, 12, BLUE, True, PP_ALIGN.CENTER)
    if i < 5: textbox(s, x+1.75, 2.43, .3, .3, "→", 18, BLUE, True, PP_ALIGN.CENTER)
textbox(s, .7, 3.65, 5.8, .35, "State giữ lại", 17, BLUE, True)
textbox(s, .7, 4.12, 5.8, 1.45, "thread_id · lesson_id · question\noptions · selected_option · trace", 18, INK)
textbox(s, 7.0, 3.65, 5.5, .35, "Guard cứng", 17, BLUE, True)
textbox(s, 7.0, 4.12, 5.5, 1.45, "Không sinh câu trả lời trước khi ý định được xác nhận. Tối đa hai vòng clarification.", 18, INK)
textbox(s, .7, 6.1, 12, .35, "CP demo: LangGraph thật · một model xuyên suốt: Gemini 3.5 Flash Lite", 12, MUTED)

# 4
s = base("Case E2E: hỏi → chọn → trả lời → bôi nguồn", "Demo", 4)
imgs = [ROOT/"artifacts/video_frames/02-interrupt.png", ROOT/"artifacts/video_frames/03-answer.png", ROOT/"artifacts/video_frames/04-highlight.png"]
labels = ["① interrupt", "② answer", "③ highlight source"]
for i, (img, label) in enumerate(zip(imgs, labels)):
    x = .55 + i * 4.2
    s.shapes.add_picture(str(img), Inches(x), Inches(1.55), width=Inches(3.85), height=Inches(2.42))
    textbox(s, x, 4.1, 3.85, .35, label, 14, BLUE, True, PP_ALIGN.CENTER)
textbox(s, .7, 4.85, 12, .9, "Dải “Xem nguồn trong bài” không mở thêm hộp dài. Khi bấm, VLearn cuộn tới và bôi sáng đúng block nguồn trong lesson.", 20, INK, True, PP_ALIGN.CENTER)
textbox(s, .7, 6.05, 12, .35, "Trace: receive_input → classify → build_options → interrupt → resume → answer", 12, MUTED, False, PP_ALIGN.CENTER)

# 5
s = base("Đo được trước khi nói là tốt", "Evaluation", 5)
rect(s, .6, 1.5, 5.5, 3.9, BLUE, BLUE, True)
textbox(s, 1.0, 1.9, 4.7, .35, "KẾT QUẢ LƯỢT 1", 11, CYAN, True, PP_ALIGN.CENTER)
textbox(s, 1.0, 2.45, 4.7, .8, "18 / 20", 44, WHITE, True, PP_ALIGN.CENTER)
textbox(s, 1.0, 3.35, 4.7, .4, "routing accuracy = 90%", 17, WHITE, True, PP_ALIGN.CENTER)
textbox(s, 1.0, 4.1, 4.7, .55, "Hard cases G01/G05/G07/G08: PASS", 13, WHITE, True, PP_ALIGN.CENTER)
textbox(s, 6.7, 1.55, 5.9, .4, "Quality bar đã chốt", 20, BLUE, True)
textbox(s, 6.7, 2.1, 5.8, .8, "≥90% routing accuracy\n+ 4 hard cases bắt buộc pass", 22, INK, True)
textbox(s, 6.7, 3.25, 5.8, .4, "UI smoke", 17, BLUE, True)
textbox(s, 6.7, 3.72, 5.8, .55, "28/28 thao tác · desktop/tablet/mobile", 18, INK)
rect(s, 6.7, 4.65, 5.6, .95, RGBColor(255,248,229), RGBColor(238,196,84), True)
textbox(s, 6.95, 4.87, 5.1, .5, "2 fail thật: “form đâu”, “link tải dữ liệu”.", 15, RGBColor(115,79,0), True)
textbox(s, .7, 6.1, 12, .35, "Golden set: eval/golden-set.json · Kết quả: eval/cp3-results.json", 12, MUTED)

# 6
s = base("Sẵn sàng cho vòng demo", "Team & next steps", 6)
people = [("Nguyễn Văn Duy", "Product · LangGraph · integration"), ("Dương Thị Ngân", "UX VLearn · responsive · validation"), ("Vũ Huy Đô", "Evidence · golden set · video")]
for i, (name, role) in enumerate(people):
    y = 1.55 + i * 1.25
    rect(s, .65, y, 5.75, .95, PALE, RGBColor(195,218,235), True)
    textbox(s, .9, y+.16, 2.5, .28, name, 15, BLUE, True)
    textbox(s, 3.15, y+.17, 2.9, .42, role, 12, MUTED)
rect(s, 7.0, 1.55, 5.65, 3.45, BLUE, BLUE, True)
textbox(s, 7.45, 1.95, 4.75, .35, "DEMO SCRIPT · 30 GIÂY", 11, CYAN, True)
textbox(s, 7.45, 2.48, 4.75, 1.8, "1. Hỏi “cho tôi link”\n2. Chọn repo nhóm\n3. Mở nguồn\n4. Bôi block trong lesson\n5. Mở trace LangGraph", 20, WHITE, True)
textbox(s, .7, 5.55, 12, .65, "Tiếp theo: cho 3 willing users dùng thử, ghi quote nguyên văn và chỉ sửa theo điểm kẹt quan sát được.", 18, INK, True)

prs.save(OUT)
print(OUT)
